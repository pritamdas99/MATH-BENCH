import glob
import json
import os
from pathlib import Path
from openpyxl import load_workbook

CROPPED_SUFFIX = "_cropped"
FINAL_OUTPUT_DIR_NAMES = ("final_model_outputs", "final_model_output")


def load_eqb_questions(workbook_path: str) -> dict[str, dict[str, str]]:
    wb = load_workbook(workbook_path, read_only=True)
    sheet = wb[wb.sheetnames[0]]

    headers = None
    question_col = None
    chapter_col = None
    id_col = None
    questions: dict[str, dict[str, str]] = {}

    for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if row_index == 1:
            headers = [str(cell).strip().lower() if cell is not None else "" for cell in row]
            for idx, name in enumerate(headers):
                if name == "id":
                    id_col = idx
                elif name == "question":
                    question_col = idx
                elif name == "chaptername":
                    chapter_col = idx
            if id_col is None or question_col is None or chapter_col is None:
                raise ValueError(
                    f"Could not find required columns 'id', 'Question', and 'ChapterName' in {workbook_path}: {headers}"
                )
            continue

        if row is None:
            continue

        eqb_id = row[id_col]
        question = row[question_col]
        chapter_name = row[chapter_col]
        if eqb_id is None or question is None or chapter_name is None:
            continue

        questions[str(eqb_id).strip()] = {
            "question": str(question),
            "ChapterName": str(chapter_name),
        }

    return questions


def find_question_and_chapter(eqb_questions: dict[str, dict[str, str]], record_id: str) -> dict[str, str] | None:
    record_id = str(record_id).strip()
    if record_id in eqb_questions:
        return eqb_questions[record_id]

    if record_id.isdigit():
        scaled = str(int(record_id) * 10)
        if scaled in eqb_questions:
            return eqb_questions[scaled]

    return None


def final_output_dir_candidates() -> list[Path]:
    bases: list[Path] = []
    for base in (Path.cwd(), Path(__file__).resolve().parent):
        bases.append(base)
        bases.extend(base.parents)

    candidates: list[Path] = []
    seen: set[Path] = set()
    for base in bases:
        for dirname in FINAL_OUTPUT_DIR_NAMES:
            for candidate in (base / dirname, base / "results" / dirname):
                candidate = candidate.resolve()
                if candidate not in seen:
                    candidates.append(candidate)
                    seen.add(candidate)

    return candidates


def find_final_output_dir() -> Path | None:
    for candidate in final_output_dir_candidates():
        if candidate.is_dir():
            return candidate
    return None


def load_final_model_predictions(final_output_dir: Path | None) -> dict[str, dict[str, str]]:
    if final_output_dir is None:
        return {}

    predictions: dict[str, dict[str, str]] = {}
    for output_path in sorted(final_output_dir.glob("final_*_results.jsonl")):
        model_key = output_path.name.removeprefix("final_").removesuffix("_results.jsonl").lower()
        predictions[model_key] = {}
        with output_path.open("r", encoding="utf-8") as infile:
            for line_number, line in enumerate(infile, start=1):
                if not line.strip():
                    continue

                record = json.loads(line)
                record_id = record.get("id")
                if record_id is None:
                    print(f"Skipping line {line_number} in {output_path}: missing 'id'")
                    continue

                predictions[model_key][str(record_id).strip()] = record.get("predicted_answer", "")

    return predictions


def predictions_for_label_file(
    jsonl_path: str,
    final_predictions: dict[str, dict[str, str]],
) -> dict[str, str]:
    file_name = Path(jsonl_path).name.lower()
    for model_key in sorted(final_predictions, key=len, reverse=True):
        if model_key in file_name:
            return final_predictions[model_key]

    return {}


def process_jsonl_file(
    jsonl_path: str,
    eqb_questions: dict[str, dict[str, str]],
    final_predictions: dict[str, dict[str, str]],
) -> tuple[int, int, int]:
    updated_count = 0
    prediction_added_count = 0
    prediction_missing_count = 0
    cropped_path = os.path.splitext(jsonl_path)[0] + f"{CROPPED_SUFFIX}.jsonl"
    predicted_answers = predictions_for_label_file(jsonl_path, final_predictions)

    with open(jsonl_path, "r", encoding="utf-8") as infile, open(cropped_path, "w", encoding="utf-8", newline="\n") as cropped_out:
        for line_number, line in enumerate(infile, start=1):
            if not line.strip():
                continue

            record = json.loads(line)
            record_id = record.get("id")
            if record_id is None:
                print(f"Skipping line {line_number} in {jsonl_path}: missing 'id'")
                continue

            question = record.get("question")
            chapter_name = record.get("ChapterName", record.get("chapterName", ""))
            if not question or not chapter_name:
                eqb_record = find_question_and_chapter(eqb_questions, record_id)
                if eqb_record is not None:
                    if not question:
                        question = eqb_record["question"]
                        updated_count += 1
                    if not chapter_name:
                        chapter_name = eqb_record["ChapterName"]
                        
            

            predicted_answer = record.get("predicted_answer", "")
            if not predicted_answer and predicted_answers:
                predicted_answer = predicted_answers.get(str(record_id).strip(), "")
                if predicted_answer:
                    prediction_added_count += 1

            if not predicted_answer and predicted_answers:
                prediction_missing_count += 1

            cropped_record = {
                "id": record_id,
                "ChapterName": chapter_name,
                "labels": record.get("label", record.get("labels", "")),
                "question": question,
                "expected_final_answer": record.get("expected_final_answer", ""),
                "predicted_answer": predicted_answer,
                "predicted_final_answer": record.get("predicted_final_answer", ""),
            }

            cropped_out.write(json.dumps(cropped_record, ensure_ascii=False) + "\n")

    return updated_count, prediction_added_count, prediction_missing_count


def main() -> None:
    cwd = os.getcwd()
    workbook_path = os.path.join(cwd, "eqb.xlsx")
    if not os.path.exists(workbook_path):
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")

    eqb_questions = load_eqb_questions(workbook_path)
    final_output_dir = find_final_output_dir()
    final_predictions = load_final_model_predictions(final_output_dir)
    if final_output_dir is None:
        print("No final model output directory found; cropped files will use existing predicted_answer values only.")
    else:
        print(f"Loaded predicted answers from {final_output_dir}")

    jsonl_files = sorted(f for f in glob.glob("*.jsonl") if not f.endswith(f"{CROPPED_SUFFIX}.jsonl"))

    if not jsonl_files:
        print("No JSONL files found in the current directory.")
        return

    total_updated = 0
    total_predictions_added = 0
    total_predictions_missing = 0
    for jsonl_path in jsonl_files:
        updated_count, prediction_added_count, prediction_missing_count = process_jsonl_file(
            jsonl_path,
            eqb_questions,
            final_predictions,
        )
        print(
            f"{jsonl_path}: created cropped JSONL, filled {updated_count} missing questions, "
            f"added {prediction_added_count} predicted answers"
        )
        if prediction_missing_count:
            print(f"{jsonl_path}: missing {prediction_missing_count} predicted answers")
        total_updated += updated_count
        total_predictions_added += prediction_added_count
        total_predictions_missing += prediction_missing_count

    print(
        f"Completed. Total questions filled: {total_updated}. "
        f"Total predicted answers added: {total_predictions_added}. "
        f"Total predicted answers missing: {total_predictions_missing}."
    )


if __name__ == "__main__":
    main()
